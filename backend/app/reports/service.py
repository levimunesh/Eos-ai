"""Sales report for a day or a period: the same KPI engine as the dashboard, so the report
can never disagree with it. Exports to Excel and PDF.
"""

import io
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any
from uuid import UUID

from sqlalchemy import select

from app.core.tenancy import OrgContext
from app.kpi import service as kpi_service
from app.kpi.definitions import quantize
from app.orgs.models import Organisation

KPI_ORDER = [
    "net_sales",
    "gross_sales",
    "discounts",
    "returns",
    "transactions",
    "avg_ticket",
    "cogs",
    "gross_profit",
    "gross_margin",
    "operating_expenses",
    "inventory_value",
]


async def build(
    ctx: OrgContext, start: date, end: date, branch_ids: list[UUID] | None
) -> dict[str, Any]:
    kpis = await kpi_service.kpis(ctx, start, end, branch_ids, compare=True)
    by_key = {k["key"]: k for k in kpis["kpis"]}
    branches = await kpi_service.breakdown(ctx, start, end, branch_ids, "branch", 100, "net")
    categories = await kpi_service.breakdown(ctx, start, end, branch_ids, "category", 20, "net")
    products = await kpi_service.breakdown(ctx, start, end, branch_ids, "product", 10, "net")
    days = (end - start).days + 1
    trend_dim = "day" if days <= 62 else "month"
    trend = await kpi_service.breakdown(ctx, start, end, branch_ids, trend_dim, 1000, "key")
    org = (
        await ctx.session.execute(select(Organisation).where(Organisation.id == ctx.org_id))
    ).scalar_one()
    partial = [k for k in kpis["kpis"] if k["status"] == "partial"]
    return {
        "organisation": org.name,
        "generated_at": datetime.now(UTC).isoformat(),
        "currency": kpis["currency"],
        "currency_decimals": kpis["currency_decimals"],
        "period": kpis["period"],
        "previous_period": kpis["previous_period"],
        "branch_ids": kpis["branch_ids"],
        "kpis": [by_key[k] for k in KPI_ORDER if k in by_key],
        "trend": {"dimension": trend_dim, "rows": trend["rows"]},
        "by_branch": branches["rows"],
        "by_category": categories["rows"],
        "top_products": products["rows"],
        "data_notes": [n for k in partial for n in k["notes"]],
        "disclaimer": (
            "Figures are computed from imported records and are not a certified "
            "accounting statement."
        ),
    }


# ---- formatting ------------------------------------------------------------------------------


def _money(v: Any, dp: int) -> str:
    if v is None:
        return "—"
    return f"{quantize(Decimal(str(v)), dp):,}"


def _pct(v: Any) -> str:
    if v is None:
        return "—"
    return f"{Decimal(str(v)) * 100:+.1f}%"


def _value(k: dict[str, Any], dp: int) -> str:
    if k["value"] is None:
        return "—"
    if k["unit"] == "count":
        return f"{int(k['value']):,}"
    if k["unit"] == "ratio":
        return f"{Decimal(str(k['value'])) * 100:.1f}%"
    return _money(k["value"], dp)


def _title(r: dict[str, Any]) -> str:
    p = r["period"]
    return (
        f"Sales report — {p['from']}"
        if p["from"] == p["to"]
        else f"Sales report — {p['from']} to {p['to']}"
    )


# ---- Excel ------------------------------------------------------------------------------------


def to_xlsx(r: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    dp = r["currency_decimals"]
    num_fmt = "#,##0" + ("." + "0" * dp if dp else "")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="F1F1EE")

    ws.append([_title(r)])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(
        [r["organisation"], f"Currency: {r['currency']}", f"Generated: {r['generated_at'][:19]}"]
    )
    ws.append([])
    ws.append(["KPI", "Value", "Previous period", "Change", "Status", "Notes"])
    for c in ws[4]:
        c.font = bold
        c.fill = head_fill
    for k in r["kpis"]:
        row = [k["label"]]
        if k["unit"] == "money":
            row += [
                Decimal(str(k["value"])) if k["value"] is not None else None,
                Decimal(str(k["previous_value"])) if k["previous_value"] is not None else None,
            ]
        elif k["unit"] == "count":
            row += [k["value"], k["previous_value"]]
        else:
            row += [
                float(k["value"]) if k["value"] is not None else None,
                float(k["previous_value"]) if k["previous_value"] is not None else None,
            ]
        row += [
            float(k["change_pct"]) if k["change_pct"] is not None else None,
            k["status"],
            "; ".join(k["notes"]),
        ]
        ws.append(row)
        n = ws.max_row
        if k["unit"] == "money":
            ws.cell(n, 2).number_format = num_fmt
            ws.cell(n, 3).number_format = num_fmt
        elif k["unit"] == "ratio":
            ws.cell(n, 2).number_format = "0.0%"
            ws.cell(n, 3).number_format = "0.0%"
        else:
            ws.cell(n, 2).number_format = "#,##0"
            ws.cell(n, 3).number_format = "#,##0"
        ws.cell(n, 4).number_format = "+0.0%;-0.0%"
    for col, width in zip("ABCDEF", (24, 18, 18, 10, 10, 60), strict=True):
        ws.column_dimensions[col].width = width
    if r["data_notes"]:
        ws.append([])
        ws.append(["Data notes"])
        ws.cell(ws.max_row, 1).font = bold
        for n in r["data_notes"]:
            ws.append([n])
    ws.append([])
    ws.append([r["disclaimer"]])

    def table(title: str, rows: list[dict[str, Any]], cols: list[tuple[str, str, str]]):
        s = wb.create_sheet(title)
        s.append([h for h, _, _ in cols])
        for c in s[1]:
            c.font = bold
            c.fill = head_fill
        for row in rows:
            out = []
            for _, key, kind in cols:
                v = row.get(key)
                if kind == "money":
                    out.append(Decimal(str(v)) if v is not None else None)
                elif kind == "int":
                    out.append(int(v) if v is not None else None)
                elif kind == "qty":
                    out.append(float(v) if v is not None else None)
                elif kind == "ratio":
                    out.append(float(v) if v is not None else None)
                else:
                    out.append(v)
            s.append(out)
        for i, (_, _, kind) in enumerate(cols, start=1):
            letter = get_column_letter(i)
            s.column_dimensions[letter].width = 28 if i == 1 else 16
            for cell in s[letter][1:]:
                if kind == "money":
                    cell.number_format = num_fmt
                elif kind == "ratio":
                    cell.number_format = "0.0%"
                elif kind in ("int", "qty"):
                    cell.number_format = "#,##0.###" if kind == "qty" else "#,##0"
                if kind != "text":
                    cell.alignment = Alignment(horizontal="right")

    cols = [
        ("Name", "label", "text"),
        (f"Net sales ({r['currency']})", "net_sales", "money"),
        ("Transactions", "transactions", "int"),
        ("Quantity", "qty", "qty"),
        (f"COGS ({r['currency']})", "cogs", "money"),
        (f"Gross profit ({r['currency']})", "gross_profit", "money"),
        ("Cost coverage", "cost_coverage_pct", "ratio"),
    ]
    table("By branch", r["by_branch"], cols)
    table("By category", r["by_category"], cols)
    table("Top products", r["top_products"], cols)
    table(
        "Trend",
        r["trend"]["rows"],
        [(r["trend"]["dimension"].title(), "label", "text"), *cols[1:3]],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- PDF --------------------------------------------------------------------------------------


def to_pdf(r: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    dp = r["currency_decimals"]
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=18, alignment=0, spaceAfter=2)
    meta = ParagraphStyle("meta", parent=styles["Normal"], fontSize=9, textColor=colors.grey)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8.5, leading=11)
    ink = colors.HexColor("#111210")
    hair = colors.HexColor("#D2D2CC")
    head_bg = colors.HexColor("#F1F1EE")

    def grid(data, col_widths, numeric_from=1):
        t = Table(data, colWidths=col_widths, repeatRows=1)
        style = [
            ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
            ("FONT", (0, 1), (-1, -1), "Helvetica", 8.5),
            ("BACKGROUND", (0, 0), (-1, 0), head_bg),
            ("TEXTCOLOR", (0, 0), (-1, -1), ink),
            ("LINEBELOW", (0, 0), (-1, 0), 0.6, hair),
            ("LINEBELOW", (0, 1), (-1, -1), 0.3, hair),
            ("ALIGN", (numeric_from, 1), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        t.setStyle(TableStyle(style))
        return t

    story = [
        Paragraph(_title(r), h1),
        Paragraph(
            f"{r['organisation']} · Currency {r['currency']} · Generated "
            f"{r['generated_at'][:16].replace('T', ' ')} UTC"
            + (
                f" · Compared with {r['previous_period']['from']} to {r['previous_period']['to']}"
                if r["previous_period"]
                else ""
            ),
            meta,
        ),
        Spacer(1, 6 * mm),
        Paragraph("Key figures", h2),
    ]
    kpi_rows = [["KPI", "Value", "Previous", "Change", "Status"]]
    for k in r["kpis"]:
        prev = dict(k, value=k["previous_value"])
        kpi_rows.append(
            [
                k["label"],
                _value(k, dp),
                _value(prev, dp),
                _pct(k["change_pct"]),
                "" if k["status"] == "complete" else k["status"],
            ]
        )
    story.append(grid(kpi_rows, [60 * mm, 32 * mm, 32 * mm, 22 * mm, 22 * mm]))

    if r["data_notes"]:
        story.append(Paragraph("Data notes", h2))
        for n in r["data_notes"]:
            story.append(Paragraph(f"• {n}", small))

    def section(title, rows, label_header):
        if not rows:
            return
        story.append(Paragraph(title, h2))
        data = [
            [
                label_header,
                f"Net sales ({r['currency']})",
                "Transactions",
                f"Gross profit ({r['currency']})",
                "Cost cov.",
            ]
        ]
        for row in rows:
            data.append(
                [
                    Paragraph(str(row["label"])[:60], small),
                    _money(row["net_sales"], dp),
                    f"{int(row['transactions']):,}",
                    _money(row["gross_profit"], dp),
                    f"{Decimal(str(row['cost_coverage_pct'])) * 100:.0f}%",
                ]
            )
        story.append(grid(data, [70 * mm, 32 * mm, 24 * mm, 32 * mm, 18 * mm]))

    section("By branch", r["by_branch"], "Branch")
    section("By category", r["by_category"], "Category")
    section("Top products", r["top_products"], "Product")
    trend_rows = r["trend"]["rows"]
    if trend_rows:
        story.append(Paragraph(f"Trend by {r['trend']['dimension']}", h2))
        data = [[r["trend"]["dimension"].title(), f"Net sales ({r['currency']})", "Transactions"]]
        for row in trend_rows:
            data.append(
                [row["label"], _money(row["net_sales"], dp), f"{int(row['transactions']):,}"]
            )
        story.append(grid(data, [50 * mm, 40 * mm, 30 * mm]))
    story += [Spacer(1, 8 * mm), Paragraph(r["disclaimer"], meta)]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=_title(r),
        author="EOS AI",
    )
    doc.build(story)
    return buf.getvalue()
