"""Phase 3 acceptance: every KPI matches hand calculations and the demo company's
independently computed answers."""

import json
from decimal import Decimal

import pytest

from tests.importing import ORDER, generate_demo, run_import

# ---------------------------------------------------------------------------------------
# A small dataset where every figure is worked out by hand below.
# ---------------------------------------------------------------------------------------
FILES = {
    "branches": "Branch Code,Branch Name\nA,Alpha\nB,Beta\n",
    "products": "Item Code,Item Description,Category\nP1,Tea,Drinks\nP2,Kettle,Home\n"
    "P3,Cake,Bakery\n",
    "fx_rates": "Rate Date,Currency,Rate To OMR\n01/01/2025,USD,0.3845\n01/01/2025,AED,0.1047\n",
    # P1 costs 0.400 OMR; P2 costs 2 USD (= 0.769 OMR); P3 has no cost.
    "product_costs": "Item Code,Effective From,Unit Cost,Currency\n01/01/2025,,,\n".replace(
        "01/01/2025,,,\n", "P1,01/01/2025,0.400,OMR\nP2,01/01/2025,2,USD\n"
    ),
    "sales": (
        "Receipt No,Branch,Date,Item Code,Qty,Unit Price,Discount,Currency,Status,Ref Receipt\n"
        "R1,A,05/01/2025 10:00,P1,2,1.000,0,OMR,COMPLETED,\n"  # net 2.000
        "R1,A,05/01/2025 10:00,P3,1,0.500,0.100,OMR,COMPLETED,\n"  # net 0.400, disc 0.100
        "R2,B,06/01/2025 11:00,P2,1,3.000,0,OMR,COMPLETED,\n"  # net 3.000
        "R3,A,07/01/2025 12:00,P1,-1,1.000,0,OMR,REFUND,R1\n"  # net -1.000
        "R4,A,07/01/2025 13:00,P1,5,1.000,0,OMR,VOID,\n"  # voided: ignored
        "R5,B,08/01/2025 09:00,P1,1,10,0,USD,COMPLETED,\n"  # 10 USD = 3.845 OMR
        "R6,B,08/01/2025 09:30,P1,1,1,0,EUR,COMPLETED,\n"  # EUR, no rate: left out
    ),
    "expenses": (
        "Date,Cost Centre,Expense Type,Amount,Currency\n"
        "10/01/2025,A,Rent,100,OMR\n"  # 100.000
        "15/01/2025,,Marketing,50,AED\n"  # 50 x 0.1047 = 5.235 (head office)
        "15/02/2025,A,Rent,100,OMR\n"  # outside the period
    ),
    "inventory_counts": (
        "Count Date,Branch,Item Code,Qty On Hand,Unit Cost\n"
        "01/01/2025,A,P1,99,0.400\n"  # superseded by the later count
        "31/01/2025,A,P1,10,0.400\n"  # 4.000
        "31/01/2025,A,P3,5,\n"  # no cost: counted, not valued
        "15/01/2025,B,P2,2,0.769\n"  # 1.538
    ),
}
PERIOD = {"from": "2025-01-05", "to": "2025-01-31"}


@pytest.fixture(scope="module")
async def hand_org(client, actor, new_org):
    owner = actor("hand")
    org = await new_org(owner, name="Eos Hand Calc", base_currency="OMR")
    for entity in (
        "branches",
        "products",
        "fx_rates",
        "product_costs",
        "sales",
        "expenses",
        "inventory_counts",
    ):
        await run_import(client, owner, org["id"], entity, f"{entity}.csv", FILES[entity].encode())
    return owner, org


def by_key(payload):
    return {k["key"]: k for k in payload["kpis"]}


async def test_hand_calculated_kpis(client, hand_org):
    owner, org = hand_org
    r = await client.get("/api/v1/kpis", params=PERIOD, headers=owner.h(org["id"]))
    assert r.status_code == 200, r.text
    k = by_key(r.json())
    expect = {
        # completed lines before discount: 2.000 + 0.500 + 3.000 + 3.845
        "gross_sales": "9.345",
        "discounts": "0.100",
        "returns": "1.000",
        "net_sales": "8.245",  # 9.345 - 0.100 - 1.000
        "transactions": 4,  # R1, R2, R3, R5 (void and unconverted EUR excluded)
        "avg_ticket": "2.061",  # 8.245 / 4 = 2.06125
        # P1: 2x0.4 - 1x0.4 + 1x0.4 = 0.800; P2: 1 x 0.769 = 0.769
        "cogs": "1.569",
        "gross_profit": "6.676",
        "gross_margin": "0.8097",  # 6.676 / 8.245
        "operating_expenses": "105.235",
        "inventory_value": "5.538",  # 4.000 + 1.538
    }
    assert {key: k[key]["value"] for key in expect} == expect
    # P3 (no cost) is 0.400 of 8.245 net: coverage = 7.845 / 8.245
    assert k["cogs"]["coverage_pct"] == "0.9515"
    assert k["gross_profit"]["status"] == "partial"
    # The EUR receipt has no rate: sales figures are flagged, not silently complete.
    assert k["net_sales"]["status"] == "partial"
    assert any("EUR" in n for n in k["net_sales"]["notes"])
    assert k["inventory_value"]["status"] == "partial"
    assert k["operating_expenses"]["status"] == "complete"
    # Previous period (1-4 Jan) had no sales: no comparison is invented.
    assert k["net_sales"]["previous_value"] is None
    assert k["net_sales"]["change_pct"] is None
    assert k["net_sales"]["last_import_at"] is not None
    assert k["net_sales"]["sources"] == ["sales"]


async def test_branch_filter_and_scope(client, hand_org, actor):
    owner, org = hand_org
    branches = {
        b["code"]: b["id"]
        for b in (await client.get("/api/v1/branches", headers=owner.h(org["id"]))).json()
    }
    r = await client.get(
        "/api/v1/kpis", params={**PERIOD, "branch_ids": branches["B"]}, headers=owner.h(org["id"])
    )
    k = by_key(r.json())
    assert k["net_sales"]["value"] == "6.845"  # 3.000 + 3.845
    assert k["transactions"]["value"] == 2
    # Head-office expenses belong to no branch, so a branch view leaves them out.
    assert k["operating_expenses"]["value"] == "0.000"

    viewer = actor("viewer-a")
    inv = await client.post(
        "/api/v1/invitations",
        json={
            "email": viewer.email,
            "role": "viewer",
            "all_branches": False,
            "branch_ids": [branches["A"]],
        },
        headers=owner.h(org["id"]),
    )
    await client.post(
        "/api/v1/invitations/accept", json={"token": inv.json()["token"]}, headers=viewer.h()
    )
    k = by_key(
        (await client.get("/api/v1/kpis", params=PERIOD, headers=viewer.h(org["id"]))).json()
    )
    assert k["net_sales"]["value"] == "1.400"  # 2.000 + 0.400 - 1.000
    assert k["transactions"]["value"] == 2
    assert k["operating_expenses"]["value"] == "100.000"
    r = await client.get(
        "/api/v1/kpis", params={**PERIOD, "branch_ids": branches["B"]}, headers=viewer.h(org["id"])
    )
    assert r.status_code == 403


async def test_period_comparison(client, hand_org):
    owner, org = hand_org
    # 7-8 Jan vs 5-6 Jan: (-1.000 + 3.845) vs (2.400 + 3.000)
    r = await client.get(
        "/api/v1/kpis",
        params={"from": "2025-01-07", "to": "2025-01-08"},
        headers=owner.h(org["id"]),
    )
    k = by_key(r.json())
    assert r.json()["previous_period"] == {"from": "2025-01-05", "to": "2025-01-06"}
    assert k["net_sales"]["value"] == "2.845"
    assert k["net_sales"]["previous_value"] == "5.400"
    assert k["net_sales"]["change_pct"] == "-0.4731"  # (2.845 - 5.4) / 5.4


async def test_breakdowns_and_drilldown(client, hand_org):
    owner, org = hand_org
    h = owner.h(org["id"])
    r = await client.get(
        "/api/v1/analytics/sales", params={**PERIOD, "dimension": "category"}, headers=h
    )
    rows = {x["key"]: x for x in r.json()["rows"]}
    assert rows["Drinks"]["net_sales"] == "4.845"  # P1: 2.000 - 1.000 + 3.845
    assert rows["Home"]["net_sales"] == "3.000"
    assert rows["Bakery"]["net_sales"] == "0.400"
    assert rows["Bakery"]["cost_coverage_pct"] == "0.0000"

    r = await client.get("/api/v1/records/sales", params={**PERIOD, "receipt": "R1"}, headers=h)
    body = r.json()
    assert body["total"] == 1
    sale = body["rows"][0]
    assert sale["net"] == "2.400" and sale["lines"] == 2
    assert sale["source"]["file_name"] == "sales.csv" and sale["source"]["row"] == 2
    d = (await client.get(f"/api/v1/records/sales/{sale['id']}", headers=h)).json()
    assert [ln["source_row"] for ln in d["lines"]] == [2, 3]


async def test_bad_ranges(client, hand_org):
    owner, org = hand_org
    h = owner.h(org["id"])
    r = await client.get(
        "/api/v1/kpis", params={"from": "2025-02-01", "to": "2025-01-01"}, headers=h
    )
    assert r.status_code == 400
    r = await client.get(
        "/api/v1/kpis", params={"from": "2015-01-01", "to": "2025-01-01"}, headers=h
    )
    assert r.status_code == 400


async def test_other_org_cannot_read_records(client, hand_org, actor, new_org):
    _, org = hand_org
    stranger = actor("x")
    await new_org(stranger)
    r = await client.get("/api/v1/kpis", params=PERIOD, headers=stranger.h(org["id"]))
    assert r.status_code == 404


# ---------------------------------------------------------------------------------------
# Demo company: compare with answers computed by the generator, independently of EOS.
# ---------------------------------------------------------------------------------------
MONEY = ["gross_sales", "discounts", "returns", "net_sales", "avg_ticket", "cogs", "gross_profit"]


@pytest.fixture(scope="module")
async def demo(client, tmp_path_factory):
    from tests.conftest import Actor

    out = generate_demo(tmp_path_factory.mktemp("demo_kpi"))
    owner = Actor("kpi-owner@eos.example")
    org = (
        await client.post(
            "/api/v1/orgs",
            json={"name": "Eos Trading Co.", "base_currency": "OMR", "timezone": "Asia/Muscat"},
            headers=owner.h(),
        )
    ).json()
    for entity, fname in ORDER:
        await run_import(client, owner, org["id"], entity, fname, (out / fname).read_bytes())
    branches = {
        b["code"]: b["id"]
        for b in (await client.get("/api/v1/branches", headers=owner.h(org["id"]))).json()
    }
    answers = json.loads((out / "expected" / "answers.json").read_text())
    return owner, org, branches, answers


def _check(k: dict, expected: dict, where: str) -> None:
    for key in MONEY:
        assert k[key]["value"] == expected[key], (where, key)
    assert k["transactions"]["value"] == expected["transactions"], (where, "transactions")
    assert k["cogs"]["coverage_pct"] == expected["cogs_coverage"], (where, "coverage")
    assert k["gross_profit"]["status"] == expected["gross_profit_status"], where
    assert k["operating_expenses"]["value"] == expected["operating_expenses"], (where, "opex")


async def _kpis(client, owner, org, start, end, branch=None):
    params = {"from": start, "to": end, "compare": "none"}
    if branch:
        params["branch_ids"] = branch
    r = await client.get("/api/v1/kpis", params=params, headers=owner.h(org["id"]))
    assert r.status_code == 200, r.text
    return by_key(r.json())


def _month_end(m: str) -> str:
    import calendar

    y, mo = map(int, m.split("-"))
    return f"{m}-{calendar.monthrange(y, mo)[1]:02d}"


async def test_demo_year_matches(client, demo):
    owner, org, _, answers = demo
    _check(await _kpis(client, owner, org, "2025-01-01", "2025-12-31"), answers["year"], "year")


async def test_demo_every_month_matches(client, demo):
    owner, org, _, answers = demo
    for m, expected in answers["by_month"].items():
        _check(await _kpis(client, owner, org, f"{m}-01", _month_end(m)), expected, m)


async def test_demo_every_branch_matches(client, demo):
    owner, org, branches, answers = demo
    for code, expected in answers["by_branch_year"].items():
        k = await _kpis(client, owner, org, "2025-01-01", "2025-12-31", branches[code])
        _check(k, expected, code)


async def test_demo_every_branch_month_matches(client, demo):
    owner, org, branches, answers = demo
    for m, per_branch in answers["by_month_branch"].items():
        for code, expected in per_branch.items():
            k = await _kpis(client, owner, org, f"{m}-01", _month_end(m), branches[code])
            _check(k, expected, f"{m} {code}")


async def test_demo_inventory_value_month_end(client, demo):
    owner, org, branches, answers = demo
    for m, per_branch in answers["inventory_value_month_end"].items():
        # Each branch must match exactly ...
        for code, value in per_branch.items():
            k = await _kpis(client, owner, org, f"{m}-01", _month_end(m), branches[code])
            assert k["inventory_value"]["value"] == value, (m, code)
        # ... and the total can differ from the sum of rounded branch values only by
        # rounding (0.0005 per branch at most).
        total = sum(Decimal(v) for v in per_branch.values())
        k = await _kpis(client, owner, org, f"{m}-01", _month_end(m))
        assert abs(Decimal(k["inventory_value"]["value"]) - total) <= Decimal("0.0015"), m


async def test_demo_expenses_by_category(client, demo):
    owner, org, _, answers = demo
    r = await client.get(
        "/api/v1/kpis",
        params={"from": "2025-01-01", "to": "2025-12-31"},
        headers=owner.h(org["id"]),
    )
    assert (
        by_key(r.json())["operating_expenses"]["value"] == (answers["year"]["operating_expenses"])
    )


# ---------------------------------------------------------------------------------------
# Reports (Phase 5): the report is the dashboard's numbers in another form.
# ---------------------------------------------------------------------------------------


async def test_report_matches_kpis_and_exports(client, hand_org, tmp_path):
    import io
    import zipfile

    owner, org = hand_org
    h = owner.h(org["id"])
    k = by_key((await client.get("/api/v1/kpis", params=PERIOD, headers=h)).json())
    r = await client.get("/api/v1/reports/sales", params=PERIOD, headers=h)
    assert r.status_code == 200, r.text
    rep = r.json()
    rk = {x["key"]: x for x in rep["kpis"]}
    for key in ("net_sales", "transactions", "gross_profit", "operating_expenses"):
        assert rk[key]["value"] == k[key]["value"], key
    assert rep["trend"]["dimension"] == "day"
    assert {row["key"] for row in rep["by_branch"]} == {"A", "B"}
    assert any("Unit cost is missing" in n for n in rep["data_notes"])

    x = await client.get(
        "/api/v1/reports/sales/export", params={**PERIOD, "format": "xlsx"}, headers=h
    )
    assert x.status_code == 200
    assert x.headers["content-disposition"].endswith('.xlsx"')
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(x.content), data_only=True)
    assert wb.sheetnames == ["Summary", "By branch", "By category", "Top products", "Trend"]
    summary = {
        row[0]: row[1]
        for row in wb["Summary"].iter_rows(min_row=5, values_only=True)
        if row and row[0]
    }
    assert str(summary["Net sales"]) == k["net_sales"]["value"]
    assert summary["Transactions"] == k["transactions"]["value"]
    assert zipfile.is_zipfile(io.BytesIO(x.content))

    p = await client.get(
        "/api/v1/reports/sales/export", params={**PERIOD, "format": "pdf"}, headers=h
    )
    assert p.status_code == 200 and p.content.startswith(b"%PDF")
    assert len(p.content) > 2000


async def test_report_respects_branch_scope(client, hand_org, actor):
    owner, org = hand_org
    branches = {
        b["code"]: b["id"]
        for b in (await client.get("/api/v1/branches", headers=owner.h(org["id"]))).json()
    }
    r = await client.get(
        "/api/v1/reports/sales",
        params={**PERIOD, "branch_ids": branches["B"]},
        headers=owner.h(org["id"]),
    )
    rep = r.json()
    assert [row["key"] for row in rep["by_branch"]] == ["B"]
    assert {x["key"]: x for x in rep["kpis"]}["net_sales"]["value"] == "6.845"


# ---------------------------------------------------------------------------------------
# Insights: computed from the same aggregates; every figure is checkable.
# ---------------------------------------------------------------------------------------


async def test_insights_on_hand_data(client, hand_org):
    owner, org = hand_org
    h = owner.h(org["id"])
    # 7-8 Jan vs 5-6 Jan: 2.845 vs 5.400 = -47.3%; branch B moved +0.845, A moved -3.400.
    r = await client.get(
        "/api/v1/insights", params={"from": "2025-01-07", "to": "2025-01-08"}, headers=h
    )
    assert r.status_code == 200, r.text
    body = r.json()
    kinds = {i["kind"]: i for i in body["insights"]}
    head = kinds["headline_change"]
    assert head["figures"]["change_pct"] == "-0.4731"
    assert (
        head["figures"]["net_sales"] == "2.845" and head["figures"]["previous_net_sales"] == "5.400"
    )
    assert head["figures"]["driver"] == "A"  # A fell by 3.400 of a 2.555 total drop
    assert "A drove" in head["text"]
    assert "missing_costs" not in kinds  # every item sold on 7-8 Jan has a cost
    assert "EUR" in kinds["missing_fx"]["text"]  # R6 has no exchange rate
    assert kinds["branch_lagging"]["figures"]["branch"] == "A"
    # Every figure quoted in the text exists in `figures` (no invented numbers).
    from app.assistant.verify import verify

    for ins in body["insights"]:
        ok, bad = verify(ins["text"], [ins["figures"]])
        assert ok, (ins["kind"], bad)


async def test_insights_no_data_period(client, hand_org):
    owner, org = hand_org
    r = await client.get(
        "/api/v1/insights",
        params={"from": "2024-01-01", "to": "2024-01-31"},
        headers=owner.h(org["id"]),
    )
    assert [i["kind"] for i in r.json()["insights"]] == ["no_data"]


async def test_insights_demo_march_ramadan(client, demo):
    owner, org, branches, answers = demo
    r = await client.get(
        "/api/v1/insights",
        params={"from": "2025-03-01", "to": "2025-03-31"},
        headers=owner.h(org["id"]),
    )
    body = r.json()
    kinds = [i["kind"] for i in body["insights"]]
    head = next(i for i in body["insights"] if i["kind"] == "headline_change")
    assert head["figures"]["net_sales"] == answers["by_month"]["2025-03"]["net_sales"]
    assert Decimal(head["figures"]["change_pct"]) > 0
    assert "missing_costs" in kinds  # demo has 5 products without a cost
    # Eid (30 Mar - 2 Apr) days run at ~55%: they should surface as unusual days.
    unusual = [i for i in body["insights"] if i["kind"] == "unusual_day"]
    assert any(i["figures"]["date"] in ("2025-03-30", "2025-03-31") for i in unusual), unusual
    for ins in body["insights"]:
        from app.assistant.verify import verify

        ok, bad = verify(ins["text"], [ins["figures"]])
        assert ok, (ins["kind"], bad)
